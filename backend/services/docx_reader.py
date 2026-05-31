"""
DOCX text extraction using python-docx.
Reads paragraphs and tables to produce clean, structured text.
"""

import io
import logging

from docx import Document

logger = logging.getLogger(__name__)


def extract_text_from_docx(file_bytes: bytes) -> str:
    """
    Extract all text from a DOCX file given as raw bytes.

    Args:
        file_bytes: Raw bytes of the uploaded DOCX file.

    Returns:
        Cleaned, concatenated text preserving paragraph structure.

    Raises:
        ValueError: If the DOCX has no extractable text.
    """
    doc = Document(io.BytesIO(file_bytes))
    text_parts: list[str] = []

    # Extract text from paragraphs
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            text_parts.append(text)

    # Extract text from tables (many resumes use tables for layout)
    for table in doc.tables:
        for row in table.rows:
            row_texts = []
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    row_texts.append(cell_text)
            if row_texts:
                text_parts.append(" | ".join(row_texts))

    full_text = "\n".join(text_parts).strip()

    if not full_text:
        raise ValueError("No text could be extracted from this DOCX file.")

    logger.info("Extracted %d characters from DOCX.", len(full_text))
    return full_text
